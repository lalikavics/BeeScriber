# -*- coding: utf-8 -*-
"""Offline tests (no GPU, no internet). Needs ffmpeg. Run:  python -m beescriber.cli selftest"""
from __future__ import annotations

import json
import shutil
import subprocess
import tempfile
import threading
import time
from pathlib import Path

from . import config, exporters, sources, textfmt, tools
from .audio import SR_BYTES, PcmBuffer
from .engine import FakeEngine, Stop
from .transcriber import Journal, TranscribeError, transcribe

FAILS: list[str] = []


def check(cond: bool, name: str) -> None:
    print(("  ok   " if cond else "  FAIL ") + name, flush=True)
    if not cond:
        FAILS.append(name)


class S(dict):
    """Minimal settings object for tests."""

    def get(self, k, d=None):
        return super().get(k, config.DEFAULTS.get(k, d))

    def output_dir(self):
        return Path(self["output_dir"])


def make_wav(path: Path, seconds: float) -> None:
    ff = tools.tool("ffmpeg")
    subprocess.run([ff, "-hide_banner", "-loglevel", "error", "-y", "-f", "lavfi", "-i",
                    f"sine=frequency=440:duration={seconds}", "-ac", "1", "-ar", "16000", str(path)], check=True)


class FakeStream:
    """Stands in for yt-dlp+ffmpeg: silence of `total` seconds from `start_at`; optionally dies at `die_at`."""
    plan: list = []

    def __init__(self, kind, source, start_at, settings):
        self.start_at = start_at
        self.buffer = PcmBuffer()
        total, die_at, self.err = FakeStream.plan.pop(0) if FakeStream.plan else (FakeStream.total, None, "")
        end = min(total, die_at) if die_at is not None else total
        self.n = max(0, int(round((end - start_at) * 16000)) * 2)

    def start(self):
        b = self.buffer

        def run():
            left = self.n
            while left > 0 and not b.closed:
                k = min(left, 320000)
                with b.cv:
                    b.buf += bytes(k)
                    b.cv.notify_all()
                left -= k
            with b.cv:
                b.eof = True
                b.cv.notify_all()
        threading.Thread(target=run, daemon=True).start()
        return self

    def errors(self):
        return self.err

    def close(self):
        self.buffer.close()


def t_parse() -> None:
    print("URL felismerés")
    p = sources.parse_input
    check(p("https://www.youtube.com/watch?v=abcDEF12345&list=PLtest0123456")["both"] is True, "video+lista link")
    check(p("https://youtube.com/playlist?list=PLtest0123456&si=x")["url"] ==
          "https://www.youtube.com/playlist?list=PLtest0123456", "lista link, rövid azonosító")
    check(p("https://youtu.be/abcDEF12345?si=abc")["video_id"] == "abcDEF12345", "youtu.be megosztás")
    check(p("https://www.youtube.com/shorts/abcDEF12345")["video_id"] == "abcDEF12345", "shorts")
    check(p("youtube.com/@SomeChannel")["url"] == "https://www.youtube.com/@SomeChannel/videos", "csatorna")
    check(p(r"C:\Videos\talk.MKV")["kind"] == "file", "helyi fájl")
    check(p("hello world")["kind"] == "invalid", "érvénytelen")
    check(sources.classify_error("ERROR: Sign in to confirm you're not a bot") == "bot", "hiba: bot")
    check(sources.classify_error("ERROR: [youtube] x: Private video") == "private", "hiba: privát")


def t_text() -> None:
    print("Szöveg")
    segs = [{"start": 0, "end": 3, "text": " Hello Dr. Smith. How are"}, {"start": 3, "end": 6, "text": " you? Fine."}]
    s = textfmt.split_sentences(segs)
    check([x["text"] for x in s] == ["Hello Dr. Smith.", "How are you?", "Fine."], "mondatbontás")
    bad = [{"start": 0, "end": 30, "text": " ".join(["word"] * 60)}]
    check(textfmt.unpunctuated_runs(bad) == 60, "írásjel nélküli rész felismerése")
    check(textfmt.unpunctuated_runs(segs) == 0, "jó szöveg nem hibás")
    check(textfmt.choose_prompt(" ".join(["word"] * 40), "en") == textfmt.style_prompt("en"), "rossz minta cseréje")
    check(textfmt.choose_prompt("This is fine. Really, it is. Yes.", "hu").startswith("This"), "jó minta marad")


def _ids(segs):
    return [s["text"] for s in segs]


def t_windows(tmp: Path) -> None:
    print("Ablakok, varratok, szünet/folytatás")
    st = S(window_seconds=300, tail_seconds=15, punct_guard=True, language="auto")
    FakeStream.total, FakeStream.plan = 720.0, []
    j = Journal(1, tmp)
    res = transcribe({"kind": "url", "source": "x", "duration": 720.0}, FakeEngine(), j, st,
                     stream_factory=FakeStream)
    texts = _ids(res["segs"])
    check(len(texts) == len(set(texts)), "nincs ismétlődés")
    check(texts[0] == "T0." and res["pos"] >= 719.5, f"teljes hossz ({res['pos']:.1f}s)")
    starts = [s["start"] for s in res["segs"]]
    check(all(b > a for a, b in zip(starts, starts[1:])), "időrend")
    gaps = [b["start"] - a["end"] for a, b in zip(res["segs"], res["segs"][1:])]
    check(max(gaps) < 0.01, "nincs kihagyott rész")
    full = texts

    # pause in the middle of window 2, then resume from the journal
    j2 = Journal(2, tmp)
    eng = FakeEngine(delay=0.01)
    calls = {"n": 0}

    def stopper():
        calls["n"] += 1
        return "pause" if 40 < calls["n"] < 10_000 and stopper.armed else None
    stopper.armed = True
    FakeStream.plan = []
    try:
        transcribe({"kind": "url", "source": "x", "duration": 720.0}, eng, j2, st, should_stop=stopper,
                   stream_factory=FakeStream)
        check(False, "szünet kivétel")
    except Stop as s:
        check(s.reason == "pause", "szünet megállít")
    pos_paused, segs_paused, _ = j2.load()
    check(0 < pos_paused < 720, f"félúton megállt ({pos_paused:.0f}s mentve)")
    stopper.armed = False
    res2 = transcribe({"kind": "url", "source": "x", "duration": 720.0}, FakeEngine(), j2, st,
                      should_stop=stopper, stream_factory=FakeStream)
    check(_ids(res2["segs"]) == full, "folytatás után azonos a teljes átirat")

    # torn journal line (crash while writing) is ignored
    with open(j2.path, "a", encoding="utf-8") as f:
        f.write('{"pos": 99')
    check(j2.load()[0] == res2["pos"], "csonka naplósor kezelése")


def t_stream_errors(tmp: Path) -> None:
    print("Megszakadó letöltés")
    st = S(window_seconds=300, tail_seconds=15, punct_guard=False, language="en")
    FakeStream.total = 700.0
    FakeStream.plan = [(700.0, 410.0, "connection reset"), (700.0, None, "")]
    import beescriber.transcriber as tr
    orig = tr.time.sleep
    tr.time.sleep = lambda *_: None
    try:
        res = transcribe({"kind": "url", "source": "x", "duration": 700.0}, FakeEngine(), Journal(3, tmp), st,
                         stream_factory=FakeStream)
        texts = _ids(res["segs"])
        check(res["pos"] >= 699.5 and len(texts) == len(set(texts)), "újracsatlakozik és folytatja")
        FakeStream.plan = [(700.0, 0.0, "ERROR: Private video")]
        try:
            transcribe({"kind": "url", "source": "x", "duration": 700.0}, FakeEngine(), Journal(4, tmp), st,
                       stream_factory=FakeStream)
            check(False, "privát videó hiba")
        except TranscribeError as e:
            check(e.code == "private", "privát videó: hibakód, nincs újrapróbálás")
    finally:
        tr.time.sleep = orig


def t_guard(tmp: Path) -> None:
    print("Írásjel-őr")
    st = S(window_seconds=300, tail_seconds=15, punct_guard=True, language="en", first_pass="plain")
    FakeStream.total, FakeStream.plan = 200.0, []
    eng = FakeEngine(unpunct=True)
    wins = []
    res = transcribe({"kind": "url", "source": "x", "duration": 200.0}, eng, Journal(5, tmp), st,
                     emit=lambda ev, **d: wins.append(1) if ev == "accepted" else None, stream_factory=FakeStream)
    check(eng.calls == 2 * len(wins) and textfmt.unpunctuated_runs(res["segs"]) == 0, "rossz ablak újra, a jobb marad")
    st = S(window_seconds=300, tail_seconds=15, punct_guard=True, language="en", first_pass="careful")
    eng = FakeEngine(unpunct=True)
    wins = []
    res = transcribe({"kind": "url", "source": "x", "duration": 200.0}, eng, Journal(6, tmp), st,
                     emit=lambda ev, **d: wins.append(1) if ev == "accepted" else None, stream_factory=FakeStream)
    check(eng.calls == len(wins) and textfmt.unpunctuated_runs(res["segs"]) == 0, "gondos első menet: nincs újrázás")
    check(len(wins) == 2, f"rövid első ablak (60 mp), utána normál ({len(wins)} ablak)")
    long_ok = [{"start": 0, "end": 30, "text": ", ".join(["one two three four five six"] * 12) + "."}]
    check(textfmt.unpunctuated_runs(long_ok) == 0, "hosszú, vesszős mondat nem hiba")
    hall = "The quick brown fox jumps over the lazy dog and wants to" + " be able to" * 40
    check(textfmt.loop_words([{"text": hall}]) > 100, "ismétlődő hurok egy szegmensen belül")
    check(textfmt.collapse_loops(hall) == "The quick brown fox jumps over the lazy dog and wants to",
          "hosszú hurok eltávolítása")
    check(textfmt.collapse_loops("No, no, no, no, I said yes.") == "No, no, no, no, I said yes.", "normál ismétlés marad")
    check(textfmt.collapse_loops("we go and we go and we go and we go and stop") == "we go and stop",
          "rövid hurok: egy példány marad")
    loop = [{"start": i, "end": i + 1, "text": "Thank you for watching the video."} for i in range(4)]
    check(textfmt.repeated_words(loop) > 0, "ismétlődő hurok felismerése")


def t_real_file(tmp: Path) -> None:
    print("Valódi ffmpeg, helyi fájl, folytatás pozícióról")
    wav = tmp / "tone.wav"
    make_wav(wav, 12.0)
    info = sources.probe_file(str(wav))
    check(abs(info["duration"] - 12.0) < 0.1, f"hossz ({info['duration']})")
    from .audio import AudioStream
    s = AudioStream("file", str(wav), 4.0).start()
    n, eof = s.buffer.wait_for(10 ** 9)
    s.close()
    check(abs(n / SR_BYTES - 8.0) < 0.05, f"kezdés 4 mp-nél ({n / SR_BYTES:.2f}s maradt)")


def t_export(tmp: Path) -> None:
    print("Kimenetek")
    segs = [{"start": 0.0, "end": 4.0, "text": " Szia, ez egy próba ő és ű betűkkel."},
            {"start": 4.0, "end": 30.0, "text": " " + " ".join(["Ez egy nagyon hosszú mondat része."] * 6)}]
    job = {"title": 'Teszt: "cím" / 1?', "video_id": "abcdefghijk", "kind": "url",
           "source": "https://www.youtube.com/watch?v=abcdefghijk", "channel": "Csatorna", "duration": 30.0,
           "lang": "hu", "playlist_id": "PL1", "playlist_title": "Lista", "pl_index": 7}
    st = S(output_dir=str(tmp / "out"), formats=list(config.FORMATS), txt_mode="paragraphs", timestamps=True,
           txt_header=True)
    out = exporters.export(job, segs, st)
    names = sorted(Path(p).name for p in out)
    check(len(out) == 6 and all(Path(p).stat().st_size > 0 for p in out), "mind a 6 formátum elkészült")
    check(names[0].startswith("007 - Teszt"), f"fájlnév sablon ({names[0]})")
    check(all("Lista" in Path(p).parent.name for p in out), "lista almappa")
    srt = Path([p for p in out if p.endswith(".srt")][0]).read_text(encoding="utf-8")
    check(srt.startswith("1\n00:00:00,000 --> 00:00:04,000") and max(len(x) for x in srt.splitlines()) <= 42,
          "SRT formátum, sorhossz")
    txt = Path([p for p in out if p.endswith(".txt")][0]).read_text(encoding="utf-8")
    check(txt.startswith('Teszt: "cím" / 1?\nCsatorna · 0:00:30 · hu\nhttps://'), "TXT címsor az elején")
    from docx import Document
    d = Document([p for p in out if p.endswith(".docx")][0])
    check(any("ő és ű" in p.text for p in d.paragraphs), "DOCX ékezetek")
    check(not list(Path(out[0]).parent.glob("*.part.*")), "nincs félkész fájl")


def t_speakers(tmp: Path) -> None:
    print("Beszélők")
    from . import diarize
    segs = [{"start": 0, "end": 4, "text": " Hello, how are you?"}, {"start": 4.2, "end": 7, "text": " Fine, thanks."},
            {"start": 7.5, "end": 9, "text": " And you?"}, {"start": 9.5, "end": 12, "text": " Great."}]
    turns = [(5.0, 9.3, "SPEAKER_07"), (0.0, 4.1, "SPEAKER_03"), (9.3, 12.5, "SPEAKER_03")]
    labels = diarize.assign(segs, turns)
    check(labels == [1, 2, 2, 1], f"beszélők hozzárendelése időátfedés alapján ({labels})")
    sp = diarize.apply(segs, labels)
    st = S(output_dir=str(tmp / "spk"), formats=["txt", "srt", "docx", "html", "pdf"], txt_mode="paragraphs",
           timestamps=False, txt_header=False, ui_lang="hu")
    job = {"title": "Interjú", "kind": "file", "source": "x.wav", "duration": 12.0}
    out = exporters.export(job, sp, st)
    txt = Path([p for p in out if p.endswith(".txt")][0]).read_text(encoding="utf-8")
    check(txt == "1. beszélő: Hello, how are you?\n\n2. beszélő: Fine, thanks. And you?\n\n1. beszélő: Great.\n",
          "TXT: új bekezdés és név minden beszélőváltásnál")
    srt = Path([p for p in out if p.endswith(".srt")][0]).read_text(encoding="utf-8")
    check("[1. beszélő] Hello" in srt and "[2. beszélő] Fine" in srt and "And you?" in srt and "] And" not in srt,
          "SRT: név csak a váltásnál")
    from docx import Document
    d = Document([p for p in out if p.endswith(".docx")][0])
    check(any(r.bold and r.text == "2. beszélő: " for p in d.paragraphs for r in p.runs), "DOCX: félkövér név")
    st2 = S(output_dir=str(tmp / "spk2"), formats=["txt"], txt_mode="paragraphs", ui_lang="en", txt_header=False)
    txt2 = Path(exporters.export(job, sp, st2)[0]).read_text(encoding="utf-8")
    check(txt2.startswith("Speaker 1: Hello"), "angol felületen: Speaker 1")
    diarize.save(99, labels, tmp)
    check(diarize.load(99, tmp) == labels, "beszélők mentése a naplóhoz")
    long = [{"start": 0.0, "end": 10.0, "text": " I think so. No, I disagree.",
             "words": [[0.0, 1.0, " I"], [1.0, 2.0, " think"], [2.0, 3.0, " so."], [5.0, 6.0, " No,"],
                       [6.0, 7.0, " I"], [7.0, 8.0, " disagree."]]}]
    split = diarize.split_by_speaker(long, [(0.0, 3.5, "A"), (4.5, 10.0, "B")])
    check([(x["speaker"], x["text"].strip()) for x in split] == [(1, "I think so."), (2, "No, I disagree.")],
          "beszélőváltás egy szegmensen belül: szétválasztás szavanként")
    check(diarize.apply(segs, split) == split, "új mentési formátum visszaolvasása")


def t_dialog_filters() -> None:
    print("Tallózó szűrők (a pywebview szabálya szerint)")
    import re
    rule = r'^([\w ]+)\((\*(?:\.(?:\w+|\*))*(?:;\*(?:\.(?:\w+|\*))*)*)\)$'   # webview.util.parse_file_type
    from . import app as appmod
    src = Path(appmod.__file__).read_text(encoding="utf-8")
    filters = re.findall(r'"([^"]+\(\*[^"]*\))"', src)
    exts = ";".join("*" + e for e in sorted(sources.MEDIA_EXT))
    filters = [f.replace("{exts}", exts) for f in filters]
    check(len(filters) >= 3 and all(re.search(rule, f) for f in filters), f"minden szűrő érvényes ({len(filters)} db)")


def t_queue(tmp: Path) -> None:
    print("Sor (SQLite) + munkaszál")
    from .jobqueue import JobQueue
    st = S(output_dir=str(tmp / "qout"), formats=["txt"], window_seconds=300, tail_seconds=15, language="en",
           punct_guard=True, pause_after_video=False)
    FakeStream.total, FakeStream.plan = 400.0, []
    events = []
    q = JobQueue(st, lambda ev, **d: events.append(ev), db_path=tmp / "q.db", jobs_dir=tmp / "jobs",
                 engine_factory=lambda status: FakeEngine(), stream_factory=FakeStream)
    items = [{"kind": "url", "source": f"https://www.youtube.com/watch?v=vid0000000{i}", "video_id": f"vid0000000{i}",
              "title": f"Video {i}", "duration": 400.0, "index": i} for i in range(1, 4)]
    ids = q.add(items, {"type": "playlist", "id": "PLx", "title": "Lista"})
    check(len(ids) == 3 and q.add(items[:1]) == [], "hozzáadás, duplikáció kiszűrve")
    q.move(ids[2], 0)
    check([j["id"] for j in q.jobs()][0] == ids[2], "átrendezés")
    q.pause_job(ids[1])
    q.start()
    t0 = time.time()
    while time.time() - t0 < 20 and (q.state()["running"] or q.current is not None):
        time.sleep(0.05)
    done = {j["id"] for j in q.jobs(history=True)}
    check(done == {ids[0], ids[2]}, "2 kész, a szüneteltetett kimaradt")
    check(q.jobs()[0]["status"] == "paused", "szüneteltetett vár")
    check("vid00000001" in "".join(q.done_keys()), "előzmény: kész videó felismerve")
    files = sorted(p.name for p in (tmp / "qout" / "Lista").glob("*.txt"))
    check(files == ["001 - Video 1.txt", "003 - Video 3.txt"], f"kimeneti fájlok {files}")
    out = q.reexport(ids[0], ["srt"])
    check(out and out[0].endswith(".srt"), "utólagos exportálás más formátumba")
    check(q.delete_history([ids[0]]) == 1 and ids[0] not in {j["id"] for j in q.jobs(history=True)}, "előzmény törlése")
    q.close()


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="beescriber-test-"))
    try:
        t_parse()
        t_text()
        t_windows(tmp)
        t_stream_errors(tmp)
        t_guard(tmp)
        if tools.tool("ffmpeg"):
            t_real_file(tmp)
        else:
            print("  (ffmpeg nincs: a valódi fájl teszt kimarad)")
        t_export(tmp)
        t_dialog_filters()
        t_speakers(tmp)
        t_queue(tmp)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print("\n" + ("MINDEN TESZT RENDBEN" if not FAILS else f"HIBÁS: {len(FAILS)} -> " + json.dumps(FAILS, ensure_ascii=False)))
    return 1 if FAILS else 0
